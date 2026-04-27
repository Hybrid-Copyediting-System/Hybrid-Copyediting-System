from __future__ import annotations

from typing import TYPE_CHECKING

from lxml.etree import _Element  # noqa: PLC2701

from ets_checker.models import CheckDetail

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument


def build_paragraph_element_index(document: DocxDocument) -> list[_Element]:
    """
    Build the same combined paragraph list that ``parser/paragraphs.iter_all``
    builds, but return the underlying lxml ``<w:p>`` elements instead of
    ``Paragraph`` models.

    Order MUST match ``parser/paragraphs.iter_all`` exactly: body paragraphs
    first, then table-cell paragraphs in walk order
    (``tables[*].rows[*].cells[*].paragraphs``). Any change to that walk order
    must be mirrored in both files.
    """
    out: list[_Element] = [p._p for p in document.paragraphs]
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    out.append(p._p)
    return out


def resolve_anchor(
    detail: CheckDetail,
    para_index: list[_Element],
) -> _Element | None:
    """
    Map a ``CheckDetail.locator`` to the ``<w:p>`` element it should attach
    to. Document-level findings — and any out-of-range or missing locator —
    fall back to the first paragraph. Returns ``None`` only if the document
    has no paragraphs at all.
    """
    if not para_index:
        return None
    loc = detail.locator
    if loc is None or loc.kind == "document":
        return para_index[0]
    if loc.kind == "paragraph":
        idx = loc.paragraph_index
        if idx is None or idx < 0 or idx >= len(para_index):
            return para_index[0]
        return para_index[idx]
    return para_index[0]
